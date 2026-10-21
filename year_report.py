"""
Year-end report: one Excel file per year for records and sharing (admins only).

  GET /admin/report/year?year=YYYY      -> Upasana_Report_YYYY.xlsx

Sheets: Summary · Upasana · By month · By zone & area · Waitlist.
Real Upasana only (trial bookings left out); member bookings and common Upasana (events).
Uses the zone/area saved on each booking, attendance (confirmed vs actual) and map links when
those columns exist, so it works on databases that do not have them yet.
"""

import io
import logging
from collections import defaultdict
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, request, send_file
from sqlalchemy import text

from admin_auth import admin_required
from model import db
import schema_check

router = Blueprint("year_report", __name__)

IST = timedelta(hours=5, minutes=30)
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
ZONES = ["A", "B", "C"]


def _col(table_col, sql, fallback="NULL"):
    return sql if schema_check.has(table_col) else fallback


def _upasana_rows(year):
    """Every real Upasana of the year: member bookings and common Upasana, oldest first."""
    zone = _col("bookings.zone_code", "COALESCE(b.zone_code, u.zone_code)", "u.zone_code")
    area = _col("bookings.host_area", "COALESCE(b.host_area, u.area)", "u.area")
    trial = "AND NOT b.is_trial" if schema_check.has("bookings.is_trial") else ""
    rows = db.session.execute(text(f"""
        SELECT b.id, b.user_id, b.booking_date, b.mahaprasad, b.booked_by_admin,
               TRIM(CONCAT_WS(' ', TRIM(u.first_name), NULLIF(TRIM(u.middle_name), ''), TRIM(u.last_name))),
               u.mobile_number, CONCAT_WS(', ', NULLIF(TRIM(u.flat_no), ''), NULLIF(TRIM(u.full_address), ''),
                                          NULLIF(TRIM(u.landmark), '')),
               u.pincode, {zone}, {area},
               {_col("bookings.map_link", "b.map_link")},
               {_col("bookings.attendance_expected", "b.attendance_expected")},
               {_col("bookings.attendance_count", "b.attendance_count")},
               {_col("bookings.attendance_note", "b.attendance_note")}
          FROM bookings b JOIN users u ON u.id = b.user_id
         WHERE b.is_active {trial} AND EXTRACT(YEAR FROM b.booking_date) = :y
    """), {"y": year}).all()
    out = [{
        "kind": "Member", "id": r[0], "user_id": r[1], "date": r[2], "mahaprasad": bool(r[3]),
        "by_admin": bool(r[4]), "host": r[5], "mobile": r[6] or "", "address": r[7] or "", "pin": r[8] or "",
        "zone": r[9] or "", "area": (r[10] or "").strip(), "map_link": r[11] or "",
        "confirmed": r[12], "actual": r[13], "note": r[14] or "", "title": "",
    } for r in rows]
    events = db.session.execute(text(f"""
        SELECT id, event_date, mahaprasad, title, host_name, contact_number, address,
               {_col("upasana_events.map_link", "map_link")},
               {_col("upasana_events.attendance_expected", "attendance_expected")},
               {_col("upasana_events.attendance_count", "attendance_count")},
               {_col("upasana_events.attendance_note", "attendance_note")}
          FROM upasana_events WHERE is_active AND EXTRACT(YEAR FROM event_date) = :y
    """), {"y": year}).all()
    out += [{
        "kind": "Common", "id": e[0], "user_id": None, "date": e[1], "mahaprasad": bool(e[2]),
        "by_admin": True, "host": e[4] or "", "mobile": e[5] or "", "address": e[6] or "", "pin": "",
        "zone": "", "area": "", "map_link": e[7] or "", "confirmed": e[8], "actual": e[9], "note": e[10] or "",
        "title": e[3] or "",
    } for e in events]
    out.sort(key=lambda x: (x["date"], x["kind"]))
    return out


def _earlier_hosting(user_ids, year):
    """{user_id: [years hosted before this year]} from real bookings."""
    if not user_ids:
        return {}
    trial = "AND NOT is_trial" if schema_check.has("bookings.is_trial") else ""
    hosted = defaultdict(list)
    for uid, yr in db.session.execute(text(
        f"SELECT user_id, EXTRACT(YEAR FROM booking_date)::int FROM bookings "
        f"WHERE is_active {trial} AND user_id = ANY(:ids) AND booking_date < make_date(:y, 1, 1) ORDER BY booking_date"
    ), {"ids": list(user_ids), "y": year}).all():
        hosted[uid].append(yr)
    return hosted


def _waitlist_rows(year):
    if not schema_check.has("upasana_waitlist.id"):
        return []
    mp = _col("upasana_waitlist.mahaprasad", "w.mahaprasad")
    rows = db.session.execute(text(f"""
        SELECT w.status, w.zone_code, TRIM(CONCAT_WS(' ', TRIM(u.first_name), TRIM(u.last_name))), u.mobile_number,
               COALESCE(NULLIF(TRIM(u.area), ''), w.host_area), w.preferred_months, {mp}, w.note, w.created_at,
               w.is_trial, w.edited, b.booking_date
          FROM upasana_waitlist w JOIN users u ON u.id = w.user_id
          LEFT JOIN bookings b ON b.id = w.allocated_booking_id AND b.is_active
         WHERE w.year = :y AND NOT w.is_trial
         ORDER BY w.created_at, w.id
    """), {"y": year}).all()
    position = defaultdict(int)
    out = []
    for r in rows:
        months = [int(m) for m in (r[5] or "").split(",") if m.strip()]
        pos = None
        if r[0] == "waiting":
            position[r[1] or "-"] += 1
            pos = position[r[1] or "-"]
        out.append({
            "status": r[0], "zone": r[1] or "", "name": r[2], "mobile": r[3] or "", "area": r[4] or "",
            "months": ", ".join(MONTHS[m - 1][:3] for m in months) or "Any month",
            "mahaprasad": "" if r[6] is None else ("Yes" if r[6] else "No"), "note": r[7] or "",
            "joined": r[8], "edited": bool(r[10]), "allocated": r[11], "position": pos,
        })
    return out


@router.route("/admin/report/year", methods=["GET"])
@admin_required
def year_report():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        year = request.args.get("year", type=int)
        if not year:
            return jsonify({"error": "year is required"}), 400
        items = _upasana_rows(year)
        if not items:
            return jsonify({"error": f"There are no Upasana in {year}."}), 404
        hosted = _earlier_hosting({x["user_id"] for x in items if x["user_id"]}, year)
        waitlist = _waitlist_rows(year)
        today = (datetime.utcnow() + IST).date()

        wb = Workbook()
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="F15700")
        title_font = Font(bold=True, size=14, color="F15700")
        bold = Font(bold=True)

        def table(ws, headers, rows, start_row=1, widths=None):
            for c, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=c, value=h)
                cell.font, cell.fill = head_font, head_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for i, row in enumerate(rows, start_row + 1):
                for c, v in enumerate(row, 1):
                    cell = ws.cell(row=i, column=c, value=v)
                    if isinstance(v, (date, datetime)):
                        cell.number_format = "DD MMM YYYY" if isinstance(v, date) and not isinstance(v, datetime) else "DD MMM YYYY HH:MM"
            ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
            for c, w in enumerate(widths or [], 1):
                ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w

        # ── Upasana ──
        ws = wb.active
        ws.title = "Upasana"
        rows = []
        for x in items:
            years = hosted.get(x["user_id"], []) if x["user_id"] else []
            rows.append([
                x["date"], WEEKDAYS[x["date"].weekday()],
                "Common Upasana" if x["kind"] == "Common" else ("Admin booking" if x["by_admin"] else "Member booking"),
                x["title"] or x["host"], x["host"] if x["title"] else "", x["mobile"], x["zone"], x["area"],
                x["address"], x["pin"], "Yes" if x["mahaprasad"] else "No",
                "" if x["kind"] == "Common" else (f"Repeat ({len(years)} earlier: {', '.join(map(str, sorted(set(years))))})" if years else "First time"),
                x["confirmed"], x["actual"], x["note"], x["map_link"],
                "Done" if x["date"] < today else ("Today" if x["date"] == today else "Upcoming"),
            ])
        table(ws, ["Date", "Day", "Type", "Upasana / Host", "Host (common)", "Mobile", "Zone", "Area / Place",
                   "Address", "PIN", "MahaPrasad", "Hosted before", "Attendance confirmed", "Actual attendance",
                   "Attendance note", "Map link", "Status"], rows,
              widths=[13, 11, 16, 28, 18, 13, 7, 22, 40, 9, 11, 26, 12, 12, 26, 30, 10])

        # ── By month ──
        ws = wb.create_sheet("By month")
        month_rows = []
        totals = defaultdict(int)
        for m in range(1, 13):
            mx = [x for x in items if x["date"].month == m]
            row = {
                "A": sum(1 for x in mx if x["zone"] == "A"), "B": sum(1 for x in mx if x["zone"] == "B"),
                "C": sum(1 for x in mx if x["zone"] == "C"), "Common": sum(1 for x in mx if x["kind"] == "Common"),
                "Total": len(mx), "Saturday": sum(1 for x in mx if x["date"].weekday() == 5),
                "Sunday": sum(1 for x in mx if x["date"].weekday() == 6),
                "MahaPrasad": sum(1 for x in mx if x["mahaprasad"]),
                "Confirmed": sum(x["confirmed"] or 0 for x in mx), "Actual": sum(x["actual"] or 0 for x in mx),
            }
            for k, v in row.items():
                totals[k] += v
            month_rows.append([MONTHS[m - 1]] + list(row.values()))
        month_rows.append(["Total"] + [totals[k] for k in ["A", "B", "C", "Common", "Total", "Saturday", "Sunday",
                                                           "MahaPrasad", "Confirmed", "Actual"]])
        table(ws, ["Month", "Zone A", "Zone B", "Zone C", "Common", "Total", "Saturdays", "Sundays",
                   "MahaPrasad", "Attendance confirmed", "Actual attendance"], month_rows,
              widths=[12, 9, 9, 9, 10, 9, 11, 10, 12, 14, 14])
        for c in range(1, 12):
            ws.cell(row=14, column=c).font = bold

        # ── By zone & area ──
        ws = wb.create_sheet("By zone & area")
        groups = defaultdict(lambda: {"n": 0, "first": 0, "confirmed": 0, "actual": 0, "mp": 0})
        for x in items:
            key = ("Common", x["title"] or "Common Upasana") if x["kind"] == "Common" else (x["zone"] or "-", x["area"] or "-")
            g = groups[key]
            g["n"] += 1
            g["mp"] += 1 if x["mahaprasad"] else 0
            g["confirmed"] += x["confirmed"] or 0
            g["actual"] += x["actual"] or 0
            if x["user_id"] and not hosted.get(x["user_id"]):
                g["first"] += 1
        zone_order = {z: i for i, z in enumerate(ZONES + ["Common", "-"])}
        table(ws, ["Zone", "Area / Place", "Upasana", "First-time hosts", "MahaPrasad", "Attendance confirmed",
                   "Actual attendance"],
              [[k[0], k[1], g["n"], g["first"], g["mp"], g["confirmed"], g["actual"]]
               for k, g in sorted(groups.items(), key=lambda kv: (zone_order.get(kv[0][0], 9), -kv[1]["n"], kv[0][1].lower()))],
              widths=[9, 30, 10, 15, 12, 14, 14])

        # ── Waitlist ──
        ws = wb.create_sheet("Waitlist")
        table(ws, ["#", "Zone", "Name", "Mobile", "Area", "Preferred months", "MahaPrasad", "Status",
                   "Date given", "Joined", "Edited", "Note"],
              [[f"#{w['position']}" if w["position"] else "", w["zone"], w["name"], w["mobile"], w["area"], w["months"],
                w["mahaprasad"], w["status"].title(), w["allocated"], w["joined"], "Yes" if w["edited"] else "",
                w["note"]] for w in waitlist],
              widths=[6, 7, 26, 13, 20, 22, 12, 11, 13, 18, 8, 30])

        # ── Summary (first sheet) ──
        ws = wb.create_sheet("Summary", 0)
        members = [x for x in items if x["kind"] == "Member"]
        common = [x for x in items if x["kind"] == "Common"]
        recorded = [x for x in items if x["actual"] is not None]
        both = [x for x in items if x["actual"] is not None and x["confirmed"] is not None]
        first_time = sum(1 for x in members if not hosted.get(x["user_id"]))
        ws["A1"] = f"Upasana year-end report {year}"
        ws["A1"].font = title_font
        ws["A2"] = f"Created {(datetime.utcnow() + IST).strftime('%d %b %Y, %I:%M %p')} IST · trial-run bookings not included"
        lines = [
            ("Upasana in total", len(items)),
            ("  Member Upasana", len(members)),
            ("    booked by members", sum(1 for x in members if not x["by_admin"])),
            ("    booked by admins (incl. waitlist dates)", sum(1 for x in members if x["by_admin"])),
            ("  Common Upasana", len(common)),
            ("On Saturdays", sum(1 for x in items if x["date"].weekday() == 5)),
            ("On Sundays", sum(1 for x in items if x["date"].weekday() == 6)),
            ("Other days", sum(1 for x in items if x["date"].weekday() not in (5, 6))),
            ("With MahaPrasad", sum(1 for x in items if x["mahaprasad"])),
            ("", None),
            ("Hosts hosting for the first time", first_time),
            ("Hosts who hosted before", len(members) - first_time),
            ("", None),
        ] + [(f"Zone {z} Upasana", sum(1 for x in members if x["zone"] == z)) for z in ZONES] + [
            ("", None),
            ("Upasana with attendance recorded", len(recorded)),
            ("Actual attendance (total)", sum(x["actual"] for x in recorded)),
            ("Average attendance per Upasana", round(sum(x["actual"] for x in recorded) / len(recorded)) if recorded else 0),
            ("Confirmed → actual (Upasana with both)", f"{sum(x['confirmed'] for x in both)} → {sum(x['actual'] for x in both)}" if both else "-"),
            ("", None),
            ("Waitlist: joined", len(waitlist)),
            ("  given a date", sum(1 for w in waitlist if w["status"] == "allocated")),
            ("  still waiting", sum(1 for w in waitlist if w["status"] == "waiting")),
            ("  withdrawn", sum(1 for w in waitlist if w["status"] == "withdrawn")),
        ]
        for i, (label, value) in enumerate(lines, 4):
            ws.cell(row=i, column=1, value=label)
            if value is not None:
                ws.cell(row=i, column=2, value=value).font = bold
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 18

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(
            buf, as_attachment=True, download_name=f"Upasana_Report_{year}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        logging.exception(f"year_report failed: {e}")
        return jsonify({"error": "Could not create the report."}), 500

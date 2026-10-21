"""
Data for the admin dashboard tabs (all endpoints need the admin login token).

  GET /admin/members/attention     members whose zone is not set, and quick-registered
                                   members who have not completed their profile
  GET /admin/booking-stats         trial-run summary and bookings per month by zone (booking year)
  GET /admin/bookings/export       booking year as an Excel file (trial bookings left out)
  GET /admin/analysis              hosts in the booking year (new vs repeat) and
                                   area-wise spread with repeat hosts

"Booking year" is the admin's booking_year setting. Trial-run bookings (bookings.is_trial)
are left out of everything except the trial summary.
"""

import io
import logging
from collections import defaultdict
from datetime import datetime, timedelta

from flask import Blueprint, jsonify, request, send_file
from sqlalchemy import extract

from model import db, Booking, User, UpasanaEvent, SundayBooking
from admin_auth import require_admin
from Booking import ZONE_LIMITS
import schema_check

router = Blueprint("admin_dashboard", __name__)

IST = timedelta(hours=5, minutes=30)
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _name(u):
    return " ".join(p for p in [u.first_name, u.middle_name, u.last_name] if p).strip()


def _member(u):
    return {
        "id": u.id,
        "name": _name(u),
        "mobile_number": u.mobile_number or "",
        "area": (u.area or "").strip(),
        "pincode": u.pincode or "",
        "zone_code": u.zone_code or "",
    }


def _booking_year():
    from upasana_settings import get_booking_settings
    return get_booking_settings()["booking_year"]


def _real_bookings(year=None):
    """Active, non-trial bookings joined with their member."""
    query = (
        db.session.query(Booking, User)
        .join(User, User.id == Booking.user_id)
        .filter(Booking.is_active == True)
    )
    if schema_check.has("bookings.is_trial"):
        query = query.filter(Booking.is_trial == False)
    if year is not None:
        query = query.filter(extract("year", Booking.booking_date) == year)
    return query.all()


# ─── User Actions ─────────────────────────────────────────────────────────────

@router.route("/admin/members/attention", methods=["GET"])
def members_attention():
    err, status, _ = require_admin()
    if err:
        return err, status
    try:
        users = User.query.order_by(User.first_name, User.last_name).all()
        zone_missing = [
            {**_member(u), "full": _full_user(u)}
            for u in users
            if (u.zone_code or "") not in ZONE_LIMITS and not u.isadmin
        ]
        profile_pending = [_member(u) for u in users if u.is_quick_registered]
        return jsonify({"zone_missing": zone_missing, "profile_pending": profile_pending}), 200
    except Exception as e:
        logging.error(f"Error in members_attention: {e}")
        return jsonify({"error": "Failed to fetch members"}), 500


def _full_user(u):
    """Same fields as GET /users, so the app can open the member's edit screen directly."""
    return {
        "id": u.id, "first_name": u.first_name, "middle_name": u.middle_name, "last_name": u.last_name,
        "email": u.email, "mobile_number": u.mobile_number, "alternate_mobile_number": u.alternate_mobile_number,
        "flat_no": u.flat_no, "full_address": u.full_address, "area": u.area, "landmark": u.landmark,
        "city": u.city, "state": u.state, "pincode": u.pincode, "anugrahit": u.anugrahit,
        "gender": u.gender, "unique_family_code": u.unique_family_code, "zone_code": u.zone_code or "",
        "created_at": u.created_at.strftime("%Y-%m-%dT%H:%M:%S+05:30")
        if schema_check.has("users.created_at") and u.created_at else None,
    }


# ─── Stats ────────────────────────────────────────────────────────────────────

@router.route("/admin/booking-stats", methods=["GET"])
def booking_stats():
    err, status, _ = require_admin()
    if err:
        return err, status
    try:
        from upasana_settings import get_trial_settings

        year = _booking_year()
        trial = get_trial_settings()

        trial_rows = (
            db.session.query(Booking, User)
            .join(User, User.id == Booking.user_id)
            .filter(Booking.is_active == True, Booking.is_trial == True)
            .all()
        ) if schema_check.has("bookings.is_trial") else []
        trial_by_zone = defaultdict(int)
        for b, u in trial_rows:
            zone = _saved_zone(b, u)
            trial_by_zone[zone if zone in ZONE_LIMITS else "Other"] += 1

        months = [{"month": m, "A": 0, "B": 0, "C": 0, "events": 0, "total": 0} for m in range(1, 13)]
        if year:
            for b, u in _real_bookings(year):
                row = months[b.booking_date.month - 1]
                zone = _saved_zone(b, u)
                if zone in ZONE_LIMITS:
                    row[zone] += 1
                row["total"] += 1
            for e in UpasanaEvent.query.filter(
                UpasanaEvent.is_active == True, extract("year", UpasanaEvent.event_date) == year
            ).all():
                months[e.event_date.month - 1]["events"] += 1
                months[e.event_date.month - 1]["total"] += 1

        return jsonify({
            "year": year,
            "trial": {
                **trial,
                "bookings": len(trial_rows),
                "members": len({b.user_id for b, _ in trial_rows}),
                "by_zone": {z: trial_by_zone.get(z, 0) for z in [*ZONE_LIMITS, "Other"]},
            },
            "per_month": months,
        }), 200
    except Exception as e:
        logging.error(f"Error in booking_stats: {e}")
        return jsonify({"error": "Failed to fetch booking stats"}), 500


@router.route("/admin/bookings/export", methods=["GET"])
def export_bookings():
    err, status, _ = require_admin()
    if err:
        return err, status
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        year = request.args.get("year", type=int) or _booking_year()
        rows = []
        for b, u in _real_bookings(year):
            rows.append({
                "Date": b.booking_date,
                "Day": WEEKDAYS[b.booking_date.weekday()],
                "Host": _name(u),
                "Mobile": u.mobile_number or "",
                "Address": ", ".join(p for p in [u.flat_no, u.full_address, u.landmark] if p),
                "Area": u.area or "",
                "PIN": u.pincode or "",
                "Zone": _saved_zone(b, u) or "",
                "Mahaprasad": "Yes" if b.mahaprasad else "No",
                "Type": "Admin booking" if b.booked_by_admin else "Member booking",
            })
        for e in UpasanaEvent.query.filter(
            UpasanaEvent.is_active == True, extract("year", UpasanaEvent.event_date) == year
        ).all():
            rows.append({
                "Date": e.event_date,
                "Day": WEEKDAYS[e.event_date.weekday()],
                "Host": e.host_name or e.title,
                "Mobile": e.contact_number or "",
                "Address": e.address or "",
                "Area": "",
                "PIN": "",
                "Zone": "",
                "Mahaprasad": "Yes" if e.mahaprasad else "No",
                "Type": f"Common event: {e.title}",
            })
        if not rows:
            return jsonify({"error": f"No bookings for {year}."}), 404

        rows.sort(key=lambda r: r["Date"])
        for r in rows:
            r["Date"] = r["Date"].strftime("%d %b %Y")
        wb = Workbook()
        ws = wb.active
        ws.title = f"Upasana {year}"
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r[h] for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="F15700")
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(width + 3, 55)
        ws.freeze_panes = "A2"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Upasana_Bookings_{year}.xlsx",
        )
    except Exception as e:
        logging.exception(f"Error in export_bookings: {e}")
        return jsonify({"error": "Failed to export bookings"}), 500


# ─── Analysis ─────────────────────────────────────────────────────────────────

def _saved_zone(b, u):
    """Zone saved on the booking, else the member's current zone (older rows)."""
    saved = b.zone_code if schema_check.has("bookings.zone_code") else None
    return saved or u.zone_code


def _saved_area(b, u):
    saved = b.host_area if schema_check.has("bookings.host_area") else None
    return (saved or u.area or "").strip()


def _hosting_by_user():
    """
    {user_id: [(year, area), ...]} — one entry per Upasana hosted, with the area it was hosted in:
    active Saturday bookings (member or admin, trial ones left out) and active Sunday bookings
    (Sunday bookings have no saved area, so the member's current area is used).
    """
    hosted = defaultdict(list)
    for b, u in _real_bookings():
        hosted[b.user_id].append((b.booking_date.year, _saved_area(b, u)))
    users = {u.id: u for u in User.query.all()}
    for sb in SundayBooking.query.filter(SundayBooking.is_active == True).all():
        u = users.get(sb.user_id)
        hosted[sb.user_id].append((sb.booking_date.year, (u.area or "").strip() if u else ""))
    return hosted


@router.route("/admin/analysis", methods=["GET"])
def analysis():
    """
    Hosts in the booking year (new vs hosted before) and the data for the area-wise table:
    every area with its member count, and every member who has hosted, with the years.
    The app filters by year (all years or one year) and works out repeat hosts from this.
    """
    err, status, _ = require_admin()
    if err:
        return err, status
    try:
        year = _booking_year()
        if not year:
            return jsonify({"error": "Booking year configuration is missing."}), 500

        hosted = _hosting_by_user()
        # Every member, including those flagged as admins (they host too)
        users = {u.id: u for u in User.query.all()}

        # Hosts in the booking year: first time vs hosted in an earlier year
        year_hosts = [uid for uid, h in hosted.items() if uid in users and any(y == year for y, _ in h)]
        repeat = sum(1 for uid in year_hosts if any(y < year for y, _ in hosted[uid]))

        # Areas as typed, grouped case-insensitively: members by their current area, plus any
        # area an Upasana was hosted in (it may no longer be anyone's current area)
        areas, labels = defaultdict(int), {}
        for u in users.values():
            key = (u.area or "").strip().lower()
            if key:
                labels.setdefault(key, (u.area or "").strip().title())
                areas[key] += 1
        for h in hosted.values():
            for _, area in h:
                key = area.lower()
                if key:
                    labels.setdefault(key, area.title())
                    areas.setdefault(key, 0)

        hosts = [
            {
                "id": uid,
                "name": _name(users[uid]),
                "mobile_number": users[uid].mobile_number or "",
                "home_key": (users[uid].area or "").strip().lower(),     # current area (for "never hosted")
                "hostings": sorted(({"year": y, "area_key": a.lower()} for y, a in h), key=lambda x: x["year"]),
            }
            for uid, h in hosted.items()
            if uid in users
        ]

        return jsonify({
            "year": year,
            "hosts": {"total": len(year_hosts), "repeat": repeat, "new": len(year_hosts) - repeat},
            "years": sorted({x["year"] for h in hosts for x in h["hostings"]}),
            "areas": [{"key": k, "area": labels[k], "members": n} for k, n in areas.items()],
            "host_list": hosts,
        }), 200
    except Exception as e:
        logging.exception(f"Error in analysis: {e}")
        return jsonify({"error": "Failed to build analysis"}), 500
